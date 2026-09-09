"""
build_excel_template.py
Generates the Pospicc data-entry Excel template (Products/Sales/Inventory/Marketing)
that excel_importer.py knows how to read. Run once to (re)generate the template:

    python build_excel_template.py
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="660B0B", end_color="660B0B", fill_type="solid")  # Pospicc brand color
LEGEND_FONT = Font(name="Arial", italic=True, size=9, color="666666")
EXAMPLE_FILL = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
BODY_FONT = Font(name="Arial", size=11)

SHEETS = {
    "Products": {
        "columns": ["SKU", "Product Name", "Category", "HPP", "Selling Price", "Vendor"],
        "legend": "Wajib diisi duluan - Sales/Inventory/Marketing merujuk ke SKU di sini. SKU harus unik & konsisten (jangan berubah walau nama produk berubah).",
        "example": ["PSP-JKT-001", "Essentiel DETIV Fitted Jacket - Candy Pink", "Jacket", 320000, 699000, "Vendor A"],
    },
    "Sales": {
        "columns": ["Order ID", "SKU", "Order Date", "Units Sold", "Revenue", "Channel"],
        "legend": "Satu baris = satu order. Order ID harus unik. SKU harus sudah ada di sheet Products. Order Date format YYYY-MM-DD. Channel contoh: website, marketplace, instagram.",
        "example": ["ORD-0001", "PSP-JKT-001", "2026-08-01", 2, 1398000, "website"],
    },
    "Inventory": {
        "columns": ["SKU", "Stock on Hand", "Reserved Stock"],
        "legend": "Satu baris = satu snapshot stock terbaru per SKU. Update sheet ini tiap kali mau refresh data stock di dashboard (misal tiap pagi).",
        "example": ["PSP-JKT-001", 6, 1],
    },
    "Marketing": {
        "columns": ["SKU", "Date", "Ads Spend", "Attributed Revenue", "Is Paid"],
        "legend": "Satu baris = satu entry performa marketing per produk per tanggal. Is Paid isi 'Yes' untuk traffic berbayar (ads), 'No' untuk organic.",
        "example": ["PSP-JKT-001", "2026-08-01", 150000, 699000, "Yes"],
    },
}

wb = Workbook()
wb.remove(wb.active)  # remove default empty sheet

for sheet_name, spec in SHEETS.items():
    ws = wb.create_sheet(sheet_name)
    columns = spec["columns"]

    # Row 1: legend/instructions (merged across all columns)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
    legend_cell = ws.cell(row=1, column=1, value=f"CARA ISI: {spec['legend']}")
    legend_cell.font = LEGEND_FONT
    legend_cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[1].height = 45

    # Row 2: headers
    for col_idx, col_name in enumerate(columns, start=1):
        cell = ws.cell(row=2, column=col_idx, value=col_name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    # Row 3: one example row (clearly marked, not real data)
    for col_idx, value in enumerate(spec["example"], start=1):
        cell = ws.cell(row=3, column=col_idx, value=value)
        cell.font = BODY_FONT
        cell.fill = EXAMPLE_FILL

    # Note next to example row
    note_col = len(columns) + 2
    note_cell = ws.cell(row=3, column=note_col, value="<- Contoh baris. Hapus/timpa dengan data asli.")
    note_cell.font = LEGEND_FONT

    # Column widths
    for col_idx, col_name in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = max(16, len(col_name) + 4)

    # Freeze header rows so they stay visible while scrolling
    ws.freeze_panes = "A4"

# Put Products first as the active sheet (it must be filled in first)
wb.move_sheet("Products", offset=-len(wb.sheetnames))
wb.active = wb.sheetnames.index("Products")

wb.save("templates/Pospicc_Data_Template.xlsx")
print("Template saved to templates/Pospicc_Data_Template.xlsx")
